#!/usr/bin/env python3
"""refresh_gallery_tags.py -- fast-refresh _diamond_demo (results page + download Excel)
from the tags CSV WITHOUT re-rendering. Run after re-tagging or applying the rule passes;
takes seconds instead of an ~11-min full rebuild. Updates the Image-tagging tab's tags,
match badges, and the Tags/Results-all download workbooks."""
import csv, re, glob, html as H
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DEMO = ROOT / "_diamond_demo"
CSV = ROOT / "Image_Recognition" / "image_recognition_results.csv"

tags_map = {r["image_file"].rsplit(".", 1)[0]: [t.strip() for t in (r.get("gemini_flash_tags") or "").split(",") if t.strip()]
            for r in csv.DictReader(open(CSV, encoding="utf-8"))}
gt = {}
xl = (glob.glob(str(Path.home() / "Downloads" / "*260628*.xlsm"))
      + glob.glob(str(Path.home() / "Library" / "CloudStorage" / "Dropbox" / "**" / "*260628*.xlsm"), recursive=True))
if xl:
    ws = load_workbook(xl[0], read_only=True, data_only=True)["Combined Test Sets"]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue
        row = list(row) + [None] * 30
        if row[4] and row[18]:
            gt[str(row[4]).rsplit(".", 1)[0].strip()] = str(row[18])


def norm(t):
    return re.sub(r"(es|s)$", "", t.strip().lower())


def score_miss(tags, pst):
    if not pst:
        return None, []
    P = [t.strip() for t in str(pst).split(",") if t.strip()]
    if not P:
        return None, []
    Fn = {norm(t) for t in tags}
    miss = [p for p in P if norm(p) not in Fn]
    return (len(P) - len(miss)) / len(P), miss


def badge(sc):
    if sc is None:
        return ""
    bg, fg = ("#fff4cc", "#8a6d00") if sc >= 0.999 else ("#e7f6ec", "#1c7a3e") if sc >= 0.75 else ("#fff0e0", "#a5590a") if sc >= 0.5 else ("#fdeaea", "#b02a2a")
    return " <span style='font:600 10px ui-monospace,Menlo,monospace;background:%s;color:%s;border-radius:999px;padding:1px 7px;margin-left:6px'>%.0f%%</span>" % (bg, fg, sc * 100)


def chips(t):
    return "".join("<span class='chip'>" + H.escape(x) + "</span>" for x in t) if t else "<span class='notags'>no Flash tags yet</span>"


def attr(t):
    return H.escape(", ".join(t)).replace("'", "&#39;")


idx = DEMO / "index.html"
h = idx.read_text(encoding="utf-8")
ts = h.find("id='tab-tags'")
ps = h.rfind("<section", 0, ts)
pe = h.find("</section>", ts)
pane = h[ps:pe]
card_re = re.compile(
    r"(<div class='design'(?:(?!<div class='design').)*?data-name=')([^']*)"
    r"('(?:(?!<div class='design').)*?data-tags=')([^']*)"
    r"('(?:(?!<div class='design').)*?<div class='dname'>)([^<]*)"
    r"((?:<span[^>]*>[^<]*</span>)?)"
    r"(</div>(?:(?!<div class='design').)*?<div class='tags'>)(.*?)(</div>)", re.DOTALL)


def repl(m):
    g1, name, g3, oldt, g5, dnm, oldb, g8, oldc, g10 = m.groups()
    stem = H.unescape(name).rsplit(".", 1)[0]
    if stem not in tags_map:
        return m.group(0)
    tags = tags_map[stem]
    sc, _ = score_miss(tags, gt.get(stem))
    return g1 + name + g3 + attr(tags) + g5 + dnm.rstrip() + badge(sc) + g8 + chips(tags) + g10


new_pane, cnt = card_re.subn(repl, pane)
idx.write_text(h[:ps] + new_pane + h[pe:], encoding="utf-8")
print(f"HTML: refreshed {cnt} tagging cards")

for xlsx in [DEMO / "downloads" / "Project_Diamond_Tags.xlsx", DEMO / "downloads" / "Project_Diamond_Results_All.xlsx"]:
    if not xlsx.exists():
        continue
    wb = load_workbook(xlsx)
    if "Image tagging" not in wb.sheetnames:
        continue
    ws = wb["Image tagging"]
    hdr = [c.value for c in ws[1]]
    col = {hh: i + 1 for i, hh in enumerate(hdr)}
    upd = 0
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, col["Design"]).value
        if not d:
            continue
        stem = str(d).rsplit(".", 1)[0]
        if stem not in tags_map:
            continue
        tags = tags_map[stem]
        ws.cell(r, col["Tags"]).value = ", ".join(tags)
        if "Match %" in col:
            sc, miss = score_miss(tags, gt.get(stem))
            c = ws.cell(r, col["Match %"])
            c.value = sc
            if sc is not None:
                c.number_format = "0%"
            if "Missed (ground-truth tags)" in col:
                ws.cell(r, col["Missed (ground-truth tags)"]).value = ", ".join(miss)
        upd += 1
    wb.save(xlsx)
    print(f"Excel: {xlsx.name} -> {upd} rows")
print("Done. Hard-refresh the page (Cmd+Shift+R).")
