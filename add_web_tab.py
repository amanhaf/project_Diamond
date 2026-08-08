#!/usr/bin/env python3
"""
add_web_tab.py
==============
Integrates the web-locator results (Image_Recognition/web_recognition_results.csv,
produced by run_locate_and_tag.py) into the EXISTING results gallery
(_diamond_demo/), to match how the other tabs behave:

  1. adds/refreshes a "Web locator" TAB in index.html (reusing the gallery's CSS),
  2. writes a standalone  downloads/Project_Diamond_Web.xlsx,
  3. adds a "Web locator" SHEET to downloads/Project_Diamond_Results_All.xlsx,
  4. adds a "Web locator" entry to the page's Download menu.

No rebuild / no re-rendering of embroidery files. Everything is idempotent, so
re-running replaces (never duplicates). run_diamond_demo.py calls this at the end
of a build (before zipping) so the tab, the Excel, and the menu entry are durable
across rebuilds and travel in the shareable zip.

    venv/bin/python add_web_tab.py
"""
import csv
import html
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "_diamond_demo"
INDEX = OUT_DIR / "index.html"
DL_DIR = OUT_DIR / "downloads"
WEB_CSV = ROOT / "Image_Recognition" / "web_recognition_results.csv"

BTN_ANCHOR = "data-tab='tags'>Image tagging</button>"
JS_ANCHOR = "forEach(function(b){"
DL_ANCHOR = "<a role='menuitem' href='downloads/Project_Diamond_Tags.xlsx' download>Image tagging</a>"

WEB_XLSX = "Project_Diamond_Web.xlsx"
COMBINED_XLSX = "Project_Diamond_Results_All.xlsx"
WEB_SHEET = "Web locator"


def _int(v):
    try:
        return int(float(v))
    except Exception:
        return 0


def _fname(url):
    return url.rstrip("/").split("/")[-1]


# ------------------------------------------------------------------ HTML tab ---
def _chips(tagstr, limit=40):
    seen, out = set(), []
    for t in (tagstr or "").split(","):
        t = t.strip()
        if t and t.lower() not in seen:
            seen.add(t.lower())
            out.append(t)
    out = out[:limit]
    if not out:
        return "<div class='notags'>no Flash tags</div>"
    chips = "".join(f"<span class='chip'>{html.escape(t)}</span>" for t in out)
    return f"<div class='tags'>{chips}</div>"


def _card(r):
    url = r.get("image_url", "")
    uid = r.get("uuid") or ""
    err = (r.get("gemini_flash_error") or "").strip()
    onerr = "this.onerror=null;this.outerHTML=\"<div class='noimg'>image blocked</div>\""
    img = (f'<a href="{html.escape(url)}" target="_blank" rel="noopener">'
           f'<img src="{html.escape(url)}" loading="lazy" onerror="{onerr}"></a>')
    body = _chips(r.get("gemini_flash_tags")) if not err else \
        f"<div class='notags'>error: {html.escape(err)[:80]}</div>"
    return (f"<div class='design'><div class='dimg'>{img}</div>"
            f"<div class='dbody'><div class='dname'>{html.escape(_fname(url))}</div>"
            f"<div class='dfmts'>{html.escape(uid)}</div>{body}</div></div>")


def build_pane(rows):
    domain = next((r["domain"] for r in rows if r.get("domain")), "")
    model = next((r["gemini_flash_model"] for r in rows if r.get("gemini_flash_model")), "gemini-2.5-flash")
    n = len(rows)
    tagged = sum(1 for r in rows if (r.get("gemini_flash_tags") or "").strip()
                 and not (r.get("gemini_flash_error") or "").strip())
    it = sum(_int(r.get("gemini_flash_input_tokens")) for r in rows)
    ot = sum(_int(r.get("gemini_flash_output_tokens")) for r in rows)
    tt = sum(_int(r.get("gemini_flash_total_tokens")) for r in rows)
    avg = (tt / n) if n else 0
    cost = (it * 0.30 + ot * 2.50) / 1_000_000

    p = ["<section id='tab-web' class='tabpane'>", "<h2>Web locator results</h2>"]
    p.append(
        f"<p class='note'>{n} design images located from "
        f"<a href='https://{html.escape(domain)}' target='_blank' rel='noopener' style='color:var(--accent)'>{html.escape(domain)}</a> "
        "by the web locator, then carried through the same pipeline as the embroidery files: each image is fetched "
        "in memory (never stored), given a content UUID with the same perceptual-hash logic, and tagged by "
        "Gemini Flash. These web images are their own records &mdash; their UUIDs are computed from the product "
        "photos, so they do not collide with the embroidery-file UUIDs in the other tabs.</p>")
    p.append(
        "<div class='tokpanel'>"
        f"<div class='tokhead'>Web tagging token usage <span class='tokmodel'>{html.escape(model)}</span></div>"
        "<div class='toktiles'>"
        f"<div class='toktile'><b>{it:,}</b><span>input tokens</span></div>"
        f"<div class='toktile'><b>{ot:,}</b><span>output tokens</span></div>"
        f"<div class='toktile'><b>{tt:,}</b><span>total tokens</span></div>"
        f"<div class='toktile'><b>{avg:,.0f}</b><span>avg / image</span></div>"
        f"<div class='toktile'><b>${cost:,.4f}</b><span>est. cost</span></div>"
        "</div>"
        f"<div class='toknote'>{n} images located &middot; {tagged} tagged &middot; "
        "billed @ $0.30/M input, $2.50/M output</div></div>")
    p.append("<div class='designs'>")
    p.extend(_card(r) for r in rows)
    p.append("</div></section>")
    return "".join(p)


def inject_html(rows):
    doc = INDEX.read_text(encoding="utf-8")
    # remove any prior web tab + menu entry (idempotent)
    doc = re.sub(r"<button class='tab-btn'[^>]*data-tab='web'>.*?</button>", "", doc, flags=re.S)
    doc = re.sub(r"<section id='tab-web' class='tabpane'>.*?</section>", "", doc, flags=re.S)
    doc = re.sub(r"<a role='menuitem' href='downloads/Project_Diamond_Web\.xlsx' download>.*?</a>", "", doc, flags=re.S)

    if BTN_ANCHOR not in doc:
        raise RuntimeError("tab-bar anchor not found")
    doc = doc.replace(BTN_ANCHOR, BTN_ANCHOR + "<button class='tab-btn' data-tab='web'>Web locator</button>", 1)

    js_idx = doc.find(JS_ANCHOR)
    if js_idx == -1:
        raise RuntimeError("tab-switch script anchor not found")
    script_idx = doc.rfind("<script>", 0, js_idx)
    if script_idx == -1:
        raise RuntimeError("<script> before tab JS not found")
    doc = doc[:script_idx] + build_pane(rows) + doc[script_idx:]

    # add the download-menu entry right after the "Image tagging" download link
    if DL_ANCHOR in doc:
        web_link = f"<a role='menuitem' href='downloads/{WEB_XLSX}' download>Web locator</a>"
        doc = doc.replace(DL_ANCHOR, DL_ANCHOR + web_link, 1)
        menu_ok = True
    else:
        menu_ok = False

    INDEX.write_text(doc, encoding="utf-8")
    return menu_ok


# ----------------------------------------------------------------- Excel -------
def _fill_web(ws, rows):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    HFILL = PatternFill("solid", fgColor="203090")
    HFONT = Font(bold=True, color="FFFFFF")
    LINKFONT = Font(color="2B7FA8", underline="single")

    ws.title = WEB_SHEET
    cols = ["Image", "Image URL", "Source page", "UUID", "Tags"]
    widths = [26, 52, 40, 38, 80]
    ws.append(cols)
    for r in rows:
        url = r.get("image_url", "")
        err = (r.get("gemini_flash_error") or "").strip()
        tags = (r.get("gemini_flash_tags") or "").strip()
        if err and not tags:
            tags = f"(error: {err})"
        ws.append([_fname(url), url, r.get("source_page", ""), r.get("uuid") or "", tags])
        c = ws.cell(row=ws.max_row, column=2)
        if url:
            c.hyperlink = url
            c.font = LINKFONT
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    for ci in range(1, len(cols) + 1):
        c = ws.cell(row=1, column=ci)
        c.fill = HFILL; c.font = HFONT; c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"


def build_excel(rows):
    from openpyxl import Workbook, load_workbook
    DL_DIR.mkdir(parents=True, exist_ok=True)

    # 1) standalone Project_Diamond_Web.xlsx
    wb = Workbook()
    _fill_web(wb.active, rows)
    wb.save(DL_DIR / WEB_XLSX)

    # 2) add/replace the "Web locator" sheet in the combined workbook
    combined = DL_DIR / COMBINED_XLSX
    if combined.exists():
        cwb = load_workbook(combined)
        if WEB_SHEET in cwb.sheetnames:
            del cwb[WEB_SHEET]
        # place it right before "Token usage" if present, else at the end
        idx = cwb.sheetnames.index("Token usage") if "Token usage" in cwb.sheetnames else len(cwb.sheetnames)
        ws = cwb.create_sheet(WEB_SHEET, idx)
        _fill_web(ws, rows)
        cwb.save(combined)
        return True
    return False


def main():
    if not INDEX.exists():
        print(f"gallery not found: {INDEX}  (run run_diamond_demo.py first)")
        sys.exit(1)
    if not WEB_CSV.exists() or WEB_CSV.stat().st_size == 0:
        print(f"no web results: {WEB_CSV}  (run run_locate_and_tag.py first)")
        sys.exit(1)

    rows = list(csv.DictReader(open(WEB_CSV, encoding="utf-8")))

    try:
        combined_ok = build_excel(rows)
    except Exception as e:
        print(f"  (Excel step skipped: {e})")
        combined_ok = False

    menu_ok = inject_html(rows)

    print(f"Added 'Web locator' tab with {len(rows)} images.")
    print(f"Excel: wrote downloads/{WEB_XLSX}" +
          (f" + 'Web locator' sheet in {COMBINED_XLSX}" if combined_ok else " (combined workbook not found yet)"))
    print(f"Download menu: {'added Web locator entry' if menu_ok else 'menu anchor not found (link not added)'}")
    print(f"Open: {INDEX}")


if __name__ == "__main__":
    main()
